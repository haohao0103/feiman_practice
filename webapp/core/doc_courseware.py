#!/usr/bin/env python3
"""文档 → 讲课逐字稿课件（指定风格）。

与视频路径的差异：
  - 没有原视频画面 → 图片全部来自「自制概念图」；但文档自带的图/表（PDF 页图、docx 内嵌图）
    会被抽出来当「原讲义图」，不浪费。
  - 文档原有标题结构（# / ##）作为大纲骨架，LLM 只负责**重排成讲课期/段**，
    不是凭空重新划分 → 保证不丢内容。
  - 时长由字数/语速估算。

支持格式：md / markdown / txt / docx / doc / rtf / html / pdf
  - docx: python-docx
  - doc / rtf / html: macOS textutil
  - pdf: workbuddy-pdf CLI（--format md 提文本，--format jpg 出页面图）
"""
import os
import re
import shutil
import subprocess
import tempfile

from . import courseware_gen as cg
from . import pdf_builder

# 仓库根目录（由文件位置推导，避免硬编码本机路径）
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))


TEXT_EXTS = {".md", ".markdown", ".txt", ".text"}
DOCX_EXTS = {".docx"}
TEXTUTIL_EXTS = {".doc", ".rtf", ".rtfd", ".html", ".htm", ".webarchive"}
PDF_EXTS = {".pdf"}
SUPPORTED = TEXT_EXTS | DOCX_EXTS | TEXTUTIL_EXTS | PDF_EXTS

_PDF_CLI = "/Users/mac/.workbuddy/external-packages/pdf-cli/workbuddy-pdf"


# ============================================================ 文本抽取
def _read_plain(path):
    for enc in ("utf-8", "gbk", "utf-16", "latin-1"):
        try:
            with open(path, encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, LookupError):
            continue
    with open(path, encoding="utf-8", errors="replace") as f:
        return f.read()


def _read_docx(path):
    try:
        import docx  # python-docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(parts)
    except Exception:
        return _read_textutil(path)


def _read_textutil(path):
    r = subprocess.run(["/usr/bin/textutil", "-convert", "txt", "-stdout", path],
                       capture_output=True, text=True, timeout=120)
    if r.returncode != 0 or not r.stdout.strip():
        raise RuntimeError(f"textutil 转换失败: {(r.stderr or '')[:300]}")
    return r.stdout


def _read_pdf(path, workdir):
    if not os.path.exists(_PDF_CLI):
        raise RuntimeError("PDF 解析需要 workbuddy-pdf CLI，未找到：" + _PDF_CLI)
    out = os.path.join(workdir, "_doc_pdf.md")
    r = subprocess.run([_PDF_CLI, path, out, "--format", "md-plain"],
                       capture_output=True, text=True, timeout=600)
    if not os.path.exists(out):
        raise RuntimeError(f"PDF 解析失败: {(r.stderr or r.stdout or '')[:400]}")
    return _read_plain(out)


def extract_text(path, workdir=None):
    """抽取文档正文（纯文本/markdown）。返回 (text, meta)。"""
    ext = os.path.splitext(path)[1].lower()
    workdir = workdir or tempfile.gettempdir()
    os.makedirs(workdir, exist_ok=True)
    if ext in TEXT_EXTS:
        text = _read_plain(path)
    elif ext in DOCX_EXTS:
        text = _read_docx(path)
    elif ext in TEXTUTIL_EXTS:
        text = _read_textutil(path)
    elif ext in PDF_EXTS:
        text = _read_pdf(path, workdir)
    else:
        raise ValueError(f"不支持的文档类型: {ext}（支持 {sorted(SUPPORTED)}）")
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 50:
        raise ValueError(f"文档内容过短（{len(text)} 字），无法生成讲课稿")
    meta = {"ext": ext, "chars": len(text),
            "words": len(re.findall(r"[\u4e00-\u9fff]", text))}
    return text, meta


# ============================================================ 文档自带图抽取
def extract_doc_images(path, out_dir, workdir=None, log=None):
    """抽取文档自带的图（PDF 每页图 / docx 内嵌图）→ frames_clean/，作为「原讲义图」。

    返回 rel_path 列表（相对 out_dir）。没有则返回空列表。
    """
    ext = os.path.splitext(path)[1].lower()
    workdir = workdir or tempfile.gettempdir()
    clean = os.path.join(out_dir, "frames_clean")
    os.makedirs(clean, exist_ok=True)
    rels = []

    if ext in PDF_EXTS and os.path.exists(_PDF_CLI):
        tmp = os.path.join(workdir, "_doc_pages")
        shutil.rmtree(tmp, ignore_errors=True)
        r = subprocess.run([_PDF_CLI, path, tmp, "--format", "jpg"],
                           capture_output=True, text=True, timeout=900)
        pages = sorted(f for f in os.listdir(tmp)) if os.path.isdir(tmp) else []
        for i, f in enumerate(pages, 1):
            if not f.lower().endswith((".jpg", ".jpeg", ".png")):
                continue
            dst = os.path.join(clean, f"V{i:03d}.jpg")
            shutil.copy2(os.path.join(tmp, f), dst)
            rels.append(f"frames_clean/V{i:03d}.jpg")
        if log:
            log(f"  文档页图抽取: {len(rels)} 张")
    elif ext in DOCX_EXTS:
        import zipfile
        try:
            with zipfile.ZipFile(path) as z:
                media = [n for n in z.namelist() if n.startswith("word/media/")]
                for i, n in enumerate(media, 1):
                    data = z.read(n)
                    if len(data) < 8000:      # 太小多半是图标/装饰
                        continue
                    dst = os.path.join(clean, f"V{i:03d}.jpg")
                    with open(dst, "wb") as f:
                        f.write(data)
                    rels.append(f"frames_clean/V{i:03d}.jpg")
            if log:
                log(f"  文档内嵌图抽取: {len(rels)} 张")
        except Exception as e:
            if log:
                log(f"  文档内嵌图抽取失败（忽略）: {e}")
    return rels


# ============================================================ 文档标题骨架
_HEAD_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")


def doc_headings(text):
    """抽取文档原有标题结构 [(层级, 标题, 行号)]，作为大纲骨架。"""
    out = []
    for i, ln in enumerate(text.split("\n"), 1):
        m = _HEAD_RE.match(ln)
        if m:
            out.append((len(m.group(1)), m.group(2).strip(), i))
    return out


# ============================================================ 大纲（文档向）
_DOC_OUTLINE_SYS = """你是一门硬核课程的内容架构师。给定一篇**文档**（讲义/教材/研报/笔记），
请把它重排成「可照着讲」的讲课大纲，输出严格 JSON。

原则：
- **文档原有结构优先**：user 会给你文档的标题骨架。保留原有标题所辖的全部内容，
  只把它们重排成「期 / 段」，不要凭空新增文档里没有的知识点，也不要漏掉任何一段。
- ★ 文档里的**非讲解内容**（标题含「承诺指标」「对账」「目录」「附录」「参考」「声明」
  「资料来源」等）**不要**作为期或段；这类内容不产生口播，交由系统放进附录。
- 按内容密度分「期」（每期一个明确主题），每期下分「段」（一段一个核心概念）。
- 每段给出：标题、start_sec/end_sec（文档无时间轴，填 0 与估算段长即可）、
  summary（本段要点）、frame_at_sec（填 0）、needs_diagram、diagram_desc。
- needs_diagram=true 的条件（宁缺毋滥）：该段讲的是一个**看文字讲不清、看图秒懂**的概念
  （结构关系/流程/对比/时序/判据表），且文档里没有现成图能替代。同一期内最多 1 张。
  diagram_desc 用中文写清这张图要表达什么、含哪些元素（节点名、箭头方向、标注文字）。

只输出 JSON：
{
  "title": "课件总标题",
  "parts": [
    {"title": "第1期 · 主题", "segments": [
      {"title": "段标题", "start_sec": 0, "end_sec": 60, "summary": "要点",
       "frame_at_sec": 0, "needs_diagram": false, "diagram_desc": ""}
    ]}
  ]
}"""


def doc_outline(cfg, text, headings, log=None):
    """文档 → 结构化大纲（保留原标题骨架）。"""
    skel = "\n".join(f"{'  ' * (lv - 1)}- {ti}（第 {i} 行）"
                     for lv, ti, i in headings) if headings else "（文档无标题结构）"
    user = (f"【文档标题骨架】\n{skel}\n\n"
            f"【文档全文】\n{text}\n\n请输出结构化大纲 JSON。")
    raw = cg.ops.call_llm(cfg, _DOC_OUTLINE_SYS, user, log=log)
    if not raw:
        raise RuntimeError("文档大纲 LLM 调用失败（双通道均不可用）")
    outline = cg._extract_json(raw)
    if not outline or "parts" not in outline:
        raise RuntimeError("文档大纲 JSON 解析失败")
    return outline


# ============================================================ 端到端
def build_doc_courseware(cfg, doc_path, out_dir, style_name="faithful", intensity=3,
                         log=None, stage_cb=None, extra_notes="", workdir=None):
    """文档 → 标杆级讲课逐字稿课件（指定风格）。

    返回 dict 含 script_path / pdf_path / corrections / outline / checks / total_images。
    """
    def _st(s, l, f=0.0):
        if stage_cb:
            try:
                stage_cb(s, l, f)
            except Exception:
                pass

    os.makedirs(out_dir, exist_ok=True)
    workdir = workdir or os.path.join(cfg.get("work_dir", "/tmp"), "doc_work")
    os.makedirs(workdir, exist_ok=True)

    _st("doc_read", "读取文档", 0.0)
    text, meta = extract_text(doc_path, workdir)
    headings = doc_headings(text)
    if log:
        log(f"  文档 {os.path.basename(doc_path)}（{meta['ext']}）"
            f"：{meta['chars']} 字符 / 中文 {meta['words']} 字 / 标题 {len(headings)} 个")

    _st("doc_read", "读取文档", 0.5)
    doc_imgs = extract_doc_images(doc_path, out_dir, workdir=workdir, log=log)
    frame_manifest = [{"key": f"V{i+1:03d}", "rel_path": r, "time_sec": 0,
                       "note": f"文档第 {i+1} 页/图", "part": "文档原图",
                       "part_idx": 0, "seg": "", "ink": 0}
                      for i, r in enumerate(doc_imgs)]
    _st("doc_read", "读取文档", 1.0)

    _st("segment_outline", "文档分段大纲", 0.0)
    outline = doc_outline(cfg, text, headings, log=log)
    if log:
        log(f"  大纲：{len(outline.get('parts', []))} 期 / "
            f"{sum(len(p.get('segments', [])) for p in outline.get('parts', []))} 段")
    _st("segment_outline", "文档分段大纲", 1.0)

    # 自制概念图（文档没有原视频画面，图全靠自制）
    diagram_specs = []
    for part in outline.get("parts", []):
        for seg in part.get("segments", []):
            if seg.get("needs_diagram") and seg.get("diagram_desc"):
                diagram_specs.append(seg["diagram_desc"])
    cap = int(cfg.get("max_diagrams_doc", 15))
    if len(diagram_specs) > cap:
        if log:
            log(f"  自制图 {len(diagram_specs)}→ 安全上限 {cap} 张（LLM 过度标记）")
        diagram_specs = diagram_specs[:cap]
    diagram_manifest = []
    if diagram_specs:
        _st("gen_diagram", "生成自制图", 0.0)
        diagram_manifest = cg._default_image_gen(diagram_specs, out_dir, log=log) or []
        _st("gen_diagram", "生成自制图", 1.0)

    _st("build_script", "生成讲课稿", 0.0)
    script = cg.generate_script(cfg, text, outline, frame_manifest, diagram_manifest,
                                log=log, style_name=style_name, intensity=intensity,
                                source_kind="document", extra_notes=extra_notes)
    truncated = cg.looks_truncated(script)
    if truncated:
        log("  ⚠ 成稿末尾不完整，疑似被 max_tokens 截断")
    _st("build_script", "生成讲课稿", 1.0)

    _st("fact_check", "核查纠错", 0.0)
    corrections = cg.fact_check(cfg, text, script, log=log)
    if log:
        log(f"  纠错：{len(corrections)} 条")
    _st("fact_check", "核查纠错", 1.0)

    est_sec = round(meta["words"] / max(60, cfg["script_speed_cpm"]) * 60, 1)
    meta_out = {"video": os.path.basename(doc_path),
                "duration_min": f"{est_sec/60:.1f}"}
    # 自检 + 质量闸门（写盘前）；不过则外科式修补
    checks, gate = cg.quality_gate(cfg, script, truncated)
    if gate:
        script2, _ = cg.repair_missing(cfg, script, checks, log=log)
        if script2 != script:
            checks2, gate2 = cg.quality_gate(cfg, script2, truncated)
            if len(gate2) < len(gate):
                script, checks, gate = script2, checks2, gate2
    _st("assemble", "组装与PDF", 0.0)
    md = cg.assemble(cfg, outline, script, frame_manifest, diagram_manifest,
                      corrections, meta_out, log=log)
    if gate:
        log("  ⚠ 质量闸门未通过：" + "；".join(gate))
        banner = ("> ⚠ **本稿未通过质量闸门**：" + "；".join(gate)
                  + "\n> 请先补齐上述项再录制（详见文末「附四 · 成稿自检」）。\n")
        lines = md.split("\n")
        md = "\n".join(lines[:2] + [banner] + lines[2:])
    else:
        log("  ✓ 质量闸门通过")
    md_path = os.path.join(out_dir, "讲课稿.md")
    with open(md_path, "w", encoding="utf-8") as fh:
        fh.write(md)
    pdf_path = os.path.join(out_dir, "讲课稿.pdf")
    try:
        pdf_builder.markdown_to_pdf(md_path, pdf_path)
        if log:
            log(f"  PDF 已生成 {pdf_path}")
    except Exception as e:
        if log:
            log(f"  PDF 生成失败（不影响 md）：{e}")
        pdf_path = None
    _st("assemble", "组装与PDF", 1.0)

    return {
        "outline": outline,
        "frame_manifest": frame_manifest,
        "diagram_manifest": diagram_manifest,
        "corrections": corrections,
        "script": script,
        "script_path": md_path,
        "pdf_path": pdf_path,
        "script_chars": len(script),
        "doc_chars": meta["chars"],
        "est_duration_min": round(est_sec / 60, 1),
        "total_images": len(frame_manifest) + len(diagram_manifest),
        "checks": checks,
        "quality_gate": gate,
    }
